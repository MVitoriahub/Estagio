from datetime import timedelta
from django.utils import timezone
from django.shortcuts import render
from django.http import HttpResponse
from django.db.models import Sum, F, DecimalField

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from venda.models import Venda, ItemVenda
from estoque.models import Estoque
from produto.models import ProdutoVariacao


PERIODOS = {
    '30':  ('Últimos 30 dias', 30),
    '60':  ('Últimos 60 dias', 60),
    '90':  ('Últimos 90 dias', 90),
    '180': ('Últimos 6 meses', 180),
    '365': ('Último 1 ano', 365),
}

TIPOS_RELATORIO = {
    'venda':           'Relatório de Venda',
    'itens-vendidos':  'Itens mais Vendidos',
    'estoque':         'Estoque Atual',
    'faturamento':     'Faturamento',
}


def relatorios_home(request):
    return render(request, 'relatorios/home.html', {'tipos': TIPOS_RELATORIO})


def escolher_periodo(request, tipo):
    if tipo not in TIPOS_RELATORIO:
        return HttpResponse('Relatório inválido', status=404)
    return render(request, 'relatorios/periodo.html', {
        'tipo': tipo,
        'tipo_label': TIPOS_RELATORIO[tipo],
        'periodos': PERIODOS,
    })


def _data_inicial(periodo):
    dias = PERIODOS.get(periodo, ('', 30))[1]
    return timezone.now() - timedelta(days=dias)


def _estilizar_cabecalho(ws, colunas):
    ws.append(colunas)
    for cel in ws[1]:
        cel.font = Font(bold=True, color='FFFFFF')
        cel.fill = PatternFill(start_color='7C3AED', end_color='7C3AED', fill_type='solid')
        cel.alignment = Alignment(horizontal='center')
    for i, col in enumerate(colunas, start=1):
        ws.column_dimensions[chr(64 + i)].width = max(18, len(col) + 4)


def gerar_relatorio(request, tipo):
    periodo = request.GET.get('periodo', '30')
    if tipo not in TIPOS_RELATORIO or periodo not in PERIODOS:
        return HttpResponse('Parâmetros inválidos', status=400)

    data_inicial = _data_inicial(periodo)
    wb = Workbook()
    ws = wb.active

    if tipo == 'venda':
        ws.title = 'Relatorio de Venda'
        _estilizar_cabecalho(ws, ['ID', 'Data', 'Cliente', 'Forma de Pagamento', 'Status', 'Valor Total (R$)'])
        vendas = (
            Venda.objects
            .filter(data_venda__gte=data_inicial, status='finalizada')
            .annotate(total=Sum(F('itens__quantidade') * F('itens__valor_unitario'), output_field=DecimalField()))
            .order_by('-data_venda')
        )
        for v in vendas:
            ws.append([
                v.pk,
                v.data_venda.strftime('%d/%m/%Y %H:%M'),
                v.cliente.nome if v.cliente else '—',
                v.get_forma_pagamento_display(),
                v.get_status_display(),
                float(v.total or 0),
            ])

    elif tipo == 'itens-vendidos':
        ws.title = 'Itens mais Vendidos'
        _estilizar_cabecalho(ws, ['Produto', 'Tamanho', 'Quantidade Vendida', 'Valor Total (R$)'])
        itens = (
            ItemVenda.objects
            .filter(venda__data_venda__gte=data_inicial, venda__status='finalizada')
            .values('variacao__produto__item', 'variacao__tamanho')
            .annotate(
                qtd=Sum('quantidade'),
                total=Sum(F('quantidade') * F('valor_unitario'), output_field=DecimalField()),
            )
            .order_by('-qtd')
        )
        for item in itens:
            ws.append([
                item['variacao__produto__item'],
                item['variacao__tamanho'],
                item['qtd'],
                float(item['total'] or 0),
            ])

    elif tipo == 'estoque':
        ws.title = 'Estoque Atual'
        _estilizar_cabecalho(ws, ['Produto', 'Tamanho', 'Quantidade em Estoque', 'Preço de Venda (R$)'])
        variacoes = ProdutoVariacao.objects.select_related('produto').order_by('produto__item', 'tamanho')
        for var in variacoes:
            ws.append([
                var.produto.item,
                var.tamanho,
                var.quantidade,
                float(var.preco_venda or 0),
            ])

        ws2 = wb.create_sheet('Movimentações no Período')
        _estilizar_cabecalho(ws2, ['Data', 'Produto', 'Tamanho', 'Tipo', 'Quantidade', 'Observação'])
        movs = (
            Estoque.objects
            .filter(data_movimentacao__gte=data_inicial)
            .select_related('variacao__produto')
            .order_by('-data_movimentacao')
        )
        for mov in movs:
            ws2.append([
                mov.data_movimentacao.strftime('%d/%m/%Y %H:%M'),
                mov.variacao.produto.item,
                mov.variacao.tamanho,
                mov.get_tipo_display(),
                mov.quantidade,
                mov.observacao,
            ])

    elif tipo == 'faturamento':
        ws.title = 'Faturamento'
        _estilizar_cabecalho(ws, ['Forma de Pagamento', 'Quantidade de Vendas', 'Valor Total (R$)'])
        vendas = (
            Venda.objects
            .filter(data_venda__gte=data_inicial, status='finalizada')
            .annotate(total=Sum(F('itens__quantidade') * F('itens__valor_unitario'), output_field=DecimalField()))
        )
        resumo = {}
        total_geral = 0
        for v in vendas:
            forma = v.get_forma_pagamento_display() or 'Não informado'
            valor = float(v.total or 0)
            resumo.setdefault(forma, {'qtd': 0, 'total': 0})
            resumo[forma]['qtd'] += 1
            resumo[forma]['total'] += valor
            total_geral += valor
        for forma, dados in resumo.items():
            ws.append([forma, dados['qtd'], dados['total']])
        ws.append(['TOTAL GERAL', vendas.count(), total_geral])
        for cel in ws[ws.max_row]:
            cel.font = Font(bold=True)

    nome_arquivo = f"{TIPOS_RELATORIO[tipo].replace(' ', '_')}_{periodo}dias.xlsx"
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{nome_arquivo}"'
    wb.save(response)
    return response


def dashboard(request):
    hoje = timezone.now()
    inicio_atual = hoje - timedelta(days=15)
    inicio_anterior = hoje - timedelta(days=30)

    def dados_periodo(data_inicio, data_fim):
        vendas = Venda.objects.filter(
            data_venda__gte=data_inicio,
            data_venda__lt=data_fim,
            status='finalizada',
        )
        qtd_vendas = vendas.count()
        faturamento = ItemVenda.objects.filter(
            venda__in=vendas
        ).aggregate(
            total=Sum(F('quantidade') * F('valor_unitario'), output_field=DecimalField())
        )['total'] or 0
        return qtd_vendas, faturamento

    qtd_atual, faturamento_atual = dados_periodo(inicio_atual, hoje)
    qtd_anterior, faturamento_anterior = dados_periodo(inicio_anterior, inicio_atual)

    def variacao(atual, anterior):
        if not anterior:
            return 100 if atual else 0
        return round(((atual - anterior) / anterior) * 100, 1)

    from clientes.models import Cliente
    clientes_novos = Cliente.objects.filter(
        criado_em__gte=inicio_atual, criado_em__lt=hoje
    ).count()
    total_clientes = Cliente.objects.count()

    # ---- Vendas Recentes (últimos itens vendidos) ----
    itens_recentes = (
        ItemVenda.objects
        .filter(venda__status='finalizada')
        .select_related('venda', 'variacao__produto')
        .order_by('-venda__data_venda')[:4]
    )
    vendas_recentes = [
        {
            'cod': item.venda.pk,
            'item': item.variacao.produto.item,
            'valor': item.quantidade * item.valor_unitario,
        }
        for item in itens_recentes
    ]

    # ---- Fluxo da Receita (faturamento por dia, últimos 15 dias) ----
    fluxo_receita = []
    fluxo_labels = []
    for i in range(14, -1, -1):
        dia = (hoje - timedelta(days=i)).date()
        total_dia = ItemVenda.objects.filter(
            venda__status='finalizada',
            venda__data_venda__date=dia,
        ).aggregate(
            total=Sum(F('quantidade') * F('valor_unitario'), output_field=DecimalField())
        )['total'] or 0
        fluxo_receita.append(float(total_dia))
        fluxo_labels.append(dia.strftime('%d/%m'))

    # ---- Mais Vendidos (top 5 produtos por quantidade, período atual) ----
    mais_vendidos_qs = (
        ItemVenda.objects
        .filter(
            venda__status='finalizada',
            venda__data_venda__gte=inicio_atual,
            venda__data_venda__lt=hoje,
        )
        .values('variacao__produto__item')
        .annotate(qtd=Sum('quantidade'))
        .order_by('-qtd')[:5]
    )
    max_qtd = mais_vendidos_qs[0]['qtd'] if mais_vendidos_qs else 1
    mais_vendidos = [
        {
            'nome': item['variacao__produto__item'],
            'qtd': item['qtd'],
            'percentual': round((item['qtd'] / max_qtd) * 100),
        }
        for item in mais_vendidos_qs
    ]

    context = {
        'total_vendas': Venda.objects.filter(status='finalizada').count(),
        'vendas_novas': qtd_atual,
        'faturamento': faturamento_atual,
        'faturamento_variacao': variacao(faturamento_atual, faturamento_anterior),
        'total_clientes': total_clientes,
        'clientes_novos': clientes_novos,
        'vendas_recentes': vendas_recentes,
        'fluxo_receita': fluxo_receita,
        'fluxo_labels': fluxo_labels,
        'mais_vendidos': mais_vendidos,
    }
    return render(request, 'relatorios/dashboard.html', context)